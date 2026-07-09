"""Excel scouting workbook generation."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .analysis import is_strong_value_signal, value_pick_signals, value_tier

HDR_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
HDR_FILL = PatternFill("solid", fgColor="1F1F1F")
HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name="Arial", size=9)
BOLD_FONT = Font(name="Arial", size=9, bold=True)
LINK_FONT = Font(name="Arial", size=9, color="0563C1", underline="single")
GREEN_BOLD = Font(name="Arial", size=9, bold=True, color="006400")
RED_FONT = Font(name="Arial", size=9, color="CC0000")
CENTER = Alignment(horizontal="center", vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)
BORDER = Border(*[Side(style="thin", color="D0D0D0")] * 4)

CAPTAIN_FILL_Y = PatternFill("solid", fgColor="92D050")
CAPTAIN_FILL_M = PatternFill("solid", fgColor="FFD966")
CAPTAIN_FONT_Y = Font(name="Arial", size=9, bold=True, color="006100")
CAPTAIN_FONT_M = Font(name="Arial", size=9, bold=True, color="806000")
ROW_TINT_CAPTAIN_Y = PatternFill("solid", fgColor="E2EFDA")
ROW_TINT_CAPTAIN_M = PatternFill("solid", fgColor="FFF2CC")
ROW_TINT_INACTIVE = PatternFill("solid", fgColor="FFF3CD")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
BLUE_FILL = PatternFill("solid", fgColor="D9E2F3")

TIER_COLORS = {"S": "FFD700", "A": "C0C0C0", "B": "CD7F32", "C": "87CEEB",
               "D": "98FB98", "E": "DDA0DD", "F": "D3D3D3"}


def _safe(val):
    """openpyxl treats strings starting with '=' as formulas; player-authored
    text (names, statements) must never be interpreted that way."""
    if isinstance(val, str) and val.startswith("="):
        return "'" + val
    return val


def _fmt_form(form):
    if not form:
        return "—"
    return f"{form['winrate']}% ({form['games']}g)"


def _sig_str(d):
    if not d["signature_heroes"]:
        return "—"
    return "\n".join(f"{s['hero']} ({s['games']}g {s['winrate']}%)"
                     for s in d["signature_heroes"])


def _last_cost_str(d):
    season = d["last_cost_season"]
    if not season:
        return "—"
    if d["was_captain_last"]:
        return f"captain ({season})"
    if d["last_cost"] is not None:
        return f"{d['last_cost']} ({season})"
    return f"undrafted ({season})"


def db_url(p):
    return f"https://www.dotabuff.com/players/{p['steam32']}"


def od_url(p):
    return f"https://www.opendota.com/players/{p['steam32']}"


def ld2l_url(p):
    return f"{config.LD2L_BASE}/profile/{p['steam64']}"


def _board_columns():
    """(header, width, getter, wrap) for the main Scouting Board sheet."""
    return [
        ("#",                     4,  lambda i, p, d: i, False),
        ("Player",               26,  lambda i, p, d: p["name"], False),
        ("Captain",               9,  lambda i, p, d: p["captain"], False),
        ("Draftable",            10,  lambda i, p, d: p["draftable"], False),
        ("Vouched",               9,  lambda i, p, d: p["vouched"], False),
        ("New?",                  7,  lambda i, p, d: "🆕" if p.get("is_new") else "", False),
        ("Listed MMR",           10,  lambda i, p, d: p["mmr"], False),
        ("OD Rank",              14,  lambda i, p, d: d["rank_str"], False),
        ("MMR Check",            20,  lambda i, p, d: d["mmr_check"], False),
        ("Est. Cost",             9,  lambda i, p, d: d["est_cost"] if d["est_cost"] is not None else "—", False),
        ("Last Cost",            18,  lambda i, p, d: _last_cost_str(d), False),
        ("Pos Prefs (1-5)",      13,  lambda i, p, d: "/".join(map(str, p["pos_prefs"])), False),
        ("Pref Pos",             10,  lambda i, p, d: p["pref_role"], False),
        ("Actual Lanes (6mo)",   24,  lambda i, p, d: d["lane_str"], False),
        ("Win%",                  7,  lambda i, p, d: d["winrate"], False),
        ("W",                     6,  lambda i, p, d: d["wins"], False),
        ("L",                     6,  lambda i, p, d: d["losses"], False),
        ("Total Games",          10,  lambda i, p, d: d["total_matches"], False),
        ("Form 30d",             12,  lambda i, p, d: _fmt_form(d["form30"]), False),
        ("Form 90d",             12,  lambda i, p, d: _fmt_form(d["form90"]), False),
        ("Last Played (days)",   13,  lambda i, p, d: d["last_match_days"], False),
        ("Recent KDA",           10,  lambda i, p, d: d["avg_kda"], False),
        ("Avg GPM",               9,  lambda i, p, d: d["avg_gpm"], False),
        ("Avg XPM",               9,  lambda i, p, d: d["avg_xpm"], False),
        ("Heroes Played",        10,  lambda i, p, d: d["versatility"], False),
        ("Top 5 Heroes",         38,  lambda i, p, d: "\n".join(d["top_heroes"]), True),
        ("Recent Heroes (last 10)", 40, lambda i, p, d: ", ".join(d["recent_heroes"][:10]), True),
        ("Signature Heroes (100g/53%+)", 38, lambda i, p, d: _sig_str(d), True),
        ("Median Lobby Rank",    14,  lambda i, p, d: d["lobby_rank_str"], False),
        ("Punches Up?",          11,  lambda i, p, d: "✅ YES" if d["punches_above"] else "", False),
        ("Solo %",                8,  lambda i, p, d: d["solo_pct"] if d["solo_pct"] is not None else "—", False),
        ("Solo WR",               9,  lambda i, p, d: d["solo_wr"] if d["solo_wr"] is not None else "—", False),
        ("Party WR",              9,  lambda i, p, d: d["party_wr"] if d["party_wr"] is not None else "—", False),
        ("League Games (6mo)",   11,  lambda i, p, d: d["league_matches"] or "—", False),
        ("League W%",            10,  lambda i, p, d: d["league_winrate"] if d["league_matches"] else "—", False),
        ("League Heroes",        35,  lambda i, p, d: ", ".join(d["league_heroes"][:8]) or "—", True),
        ("Value Tier",           20,  lambda i, p, d: value_tier(p["mmr"], d), False),
        ("Your Rating",          10,  lambda i, p, d: "", False),
        ("Draft Target Rd",      12,  lambda i, p, d: "", False),
        ("Notes",                30,  lambda i, p, d: "", True),
        ("Dotabuff",             38,  lambda i, p, d: db_url(p), False),
        ("OpenDota",             38,  lambda i, p, d: od_url(p), False),
        ("LD2L",                 30,  lambda i, p, d: ld2l_url(p), False),
        ("Statement",            50,  lambda i, p, d: p["statement"], True),
    ]


def _build_board(ws, sorted_players):
    cols = _board_columns()
    headers = [c[0] for c in cols]
    cix = {h: n for n, h in enumerate(headers, 1)}  # header -> column index

    for n, (header, width, _, _) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=n, value=header)
        cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, HDR_ALIGN, BORDER
        ws.column_dimensions[get_column_letter(n)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(sorted_players) + 1}"

    for row, pd in enumerate(sorted_players, 2):
        p, d = pd["player"], pd["data"]
        inactive = d["last_match_days"] is not None and d["last_match_days"] > config.INACTIVE_DAYS

        # Row base tint: captain status wins over the inactive warning
        row_tint = None
        if p["captain"] == "Y":
            row_tint = ROW_TINT_CAPTAIN_Y
        elif p["captain"] == "M":
            row_tint = ROW_TINT_CAPTAIN_M
        elif inactive:
            row_tint = ROW_TINT_INACTIVE

        for n, (header, _, get, wrap) in enumerate(cols, 1):
            cell = ws.cell(row=row, column=n, value=_safe(get(row - 1, p, d)))
            cell.font = DATA_FONT
            cell.border = BORDER
            cell.alignment = WRAP if wrap else CENTER
            if row_tint is not None:
                cell.fill = row_tint

        def cell_at(header):
            return ws.cell(row=row, column=cix[header])

        if p["captain"] == "Y":
            cell_at("Captain").fill, cell_at("Captain").font = CAPTAIN_FILL_Y, CAPTAIN_FONT_Y
        elif p["captain"] == "M":
            cell_at("Captain").fill, cell_at("Captain").font = CAPTAIN_FILL_M, CAPTAIN_FONT_M

        if p.get("is_new"):
            cell_at("New?").fill = GREEN_FILL

        if inactive:
            cell_at("Last Played (days)").font = Font(name="Arial", size=9, bold=True, color="CC0000")

        if d["winrate"] >= 53:
            cell_at("Win%").font = GREEN_BOLD
        elif d["winrate"] < 48 and d["total_matches"] > 100:
            cell_at("Win%").font = RED_FONT

        if d["mmr_check"].startswith("⚠"):
            cell_at("MMR Check").font = GREEN_BOLD
            cell_at("MMR Check").fill = GREEN_FILL
        elif d["mmr_check"]:
            cell_at("MMR Check").font = RED_FONT

        if d["punches_above"]:
            cell_at("Punches Up?").fill = GREEN_FILL
            cell_at("Punches Up?").font = GREEN_BOLD
            cell_at("Median Lobby Rank").font = GREEN_BOLD

        if d["league_matches"]:
            cell_at("League Games (6mo)").fill = BLUE_FILL
            cell_at("League Games (6mo)").font = BOLD_FONT

        tier = cell_at("Value Tier").value or ""
        if tier[:1] in TIER_COLORS:
            cell_at("Value Tier").fill = PatternFill("solid", fgColor=TIER_COLORS[tier[0]])

        for header, url in (("Dotabuff", db_url(p)), ("OpenDota", od_url(p)), ("LD2L", ld2l_url(p))):
            cell_at(header).font = LINK_FONT
            cell_at(header).hyperlink = url


def _build_by_role(ws, sorted_players):
    ws.sheet_properties.tabColor = "4472C4"
    pos_names = {1: "Pos 1 (Carry)", 2: "Pos 2 (Mid)", 3: "Pos 3 (Offlane)",
                 4: "Pos 4 (Soft Support)", 5: "Pos 5 (Hard Support)"}

    def plays_pos(p, pos):
        # rated 1-2, or it's their best-rated position (covers unfilled/odd forms)
        return p["pos_prefs"][pos - 1] <= 2 or p["pos_prefs"][pos - 1] == min(p["pos_prefs"])

    groups = [(pos_names[pos],
               [pd for pd in sorted_players if plays_pos(pd["player"], pos)])
              for pos in range(1, 6)]
    groups.append(("Flexible (all positions)",
                   [pd for pd in sorted_players if pd["player"]["pref_role"] == "Any"]))

    row = 1
    for group_name, group in groups:
        cell = ws.cell(row=row, column=1, value=f"{group_name} ({len(group)} players)")
        cell.font = Font(bold=True, size=12, name="Arial", color="FFFFFF")
        for c in range(1, 11):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="2D2D2D")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        row += 1

        sub = ["Name", "MMR", "Captain", "Rank", "Win%", "Form 30d", "KDA", "GPM",
               "Last Played", "Top Heroes"]
        for c, h in enumerate(sub, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = Font(bold=True, name="Arial", size=9)
            cell.fill = BLUE_FILL
        row += 1

        for pd in group:
            p, d = pd["player"], pd["data"]
            cap_cell = ws.cell(row=row, column=3, value=p["captain"])
            if p["captain"] == "Y":
                cap_cell.fill, cap_cell.font = CAPTAIN_FILL_Y, CAPTAIN_FONT_Y
            elif p["captain"] == "M":
                cap_cell.fill, cap_cell.font = CAPTAIN_FILL_M, CAPTAIN_FONT_M
            else:
                cap_cell.font = DATA_FONT
            lp = f"{d['last_match_days']}d ago" if d["last_match_days"] is not None else "?"
            vals = [p["name"], p["mmr"], None, d["rank_str"], d["winrate"],
                    _fmt_form(d["form30"]), d["avg_kda"], d["avg_gpm"], lp,
                    ", ".join(d["top_heroes"][:3])]
            for c, v in enumerate(vals, 1):
                if v is None:
                    continue
                ws.cell(row=row, column=c, value=_safe(v)).font = DATA_FONT
            row += 1
        row += 1

    for c, w in {"A": 26, "B": 8, "C": 9, "D": 14, "E": 7, "F": 12, "G": 7,
                 "H": 7, "I": 12, "J": 55}.items():
        ws.column_dimensions[c].width = w


def _build_value_picks(ws, sorted_players):
    ws.sheet_properties.tabColor = "FF6347"
    ws.cell(row=1, column=1, value="🔥 POTENTIAL VALUE PICKS - Players Who May Outperform Their MMR").font = \
        Font(bold=True, size=13, name="Arial")
    ws.merge_cells("A1:H1")

    headers = ["Player", "MMR", "Rank", "Win%", "Total Games", "KDA",
               "Why They're Interesting", "Risk"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill, cell.alignment = HDR_FONT, PatternFill("solid", fgColor="CC3300"), HDR_ALIGN

    picks = []
    for pd in sorted_players:
        reasons, risks = value_pick_signals(pd["player"], pd["data"])
        if len(reasons) >= 2 or (reasons and is_strong_value_signal(pd["data"])):
            picks.append((pd, reasons, risks))

    row = 4
    for pd, reasons, risks in sorted(picks, key=lambda x: x[0]["player"]["mmr"]):
        p, d = pd["player"], pd["data"]
        vals = [p["name"], p["mmr"], d["rank_str"], d["winrate"], d["total_matches"],
                d["avg_kda"], " | ".join(reasons), " | ".join(risks) if risks else "Low risk"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=_safe(v))
            cell.font = BOLD_FONT if c == 1 else DATA_FONT
            cell.border = BORDER
            cell.fill = PatternFill("solid", fgColor="FFF0E0")
            if c >= 7:
                cell.alignment = WRAP
        ws.row_dimensions[row].height = 40
        row += 1

    for c, w in {"A": 26, "B": 8, "C": 14, "D": 7, "E": 10, "F": 7, "G": 65, "H": 35}.items():
        ws.column_dimensions[c].width = w


def _build_signatures(ws, sorted_players):
    ws.sheet_properties.tabColor = "7030A0"
    ws.cell(row=1, column=1, value=f"🎯 SIGNATURE HEROES — {config.SIGNATURE_MIN_GAMES}+ Games & "
            f"{config.SIGNATURE_MIN_WINRATE:g}%+ Win Rate").font = Font(bold=True, size=13, name="Arial")
    ws.merge_cells("A1:G1")
    ws.cell(row=2, column=1, value="Heroes players can reliably perform on. Target these in draft "
            "or ban them against opponents.").font = Font(italic=True, name="Arial", size=9, color="666666")
    ws.merge_cells("A2:G2")

    headers = ["Player", "MMR", "Hero", "Games", "Wins", "Win Rate", "Pref Pos"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = \
            HDR_FONT, PatternFill("solid", fgColor="4A1A7A"), HDR_ALIGN, BORDER

    all_sigs = [(pd["player"], sh)
                for pd in sorted_players for sh in pd["data"]["signature_heroes"]]
    all_sigs.sort(key=lambda x: (-x[1]["winrate"], -x[1]["games"]))

    row = 5
    for p, sh in all_sigs:
        wr = sh["winrate"]
        vals = [p["name"], p["mmr"], sh["hero"], sh["games"], sh["wins"], f"{wr}%", p["pref_role"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=_safe(v))
            cell.border = BORDER
            cell.alignment = CENTER if c > 1 else WRAP
            cell.font = DATA_FONT
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3E8FF")
        ws.cell(row=row, column=1).font = BOLD_FONT
        ws.cell(row=row, column=3).font = Font(name="Arial", size=10, bold=True, color="4A1A7A")
        wr_cell = ws.cell(row=row, column=6)
        if wr >= 60:
            wr_cell.font, wr_cell.fill = GREEN_BOLD, GREEN_FILL
        elif wr >= 55:
            wr_cell.font = GREEN_BOLD
        row += 1

    row += 1
    ws.cell(row=row, column=1, value=f"Total signature heroes found: {len(all_sigs)}").font = BOLD_FONT
    players_with = len({p["name"] for p, _ in all_sigs})
    ws.cell(row=row + 1, column=1,
            value=f"Players with at least one: {players_with}/{len(sorted_players)}").font = DATA_FONT

    for c, w in {"A": 28, "B": 8, "C": 22, "D": 8, "E": 8, "F": 10, "G": 16}.items():
        ws.column_dimensions[c].width = w


def _build_lobby_intel(ws, sorted_players):
    ws.sheet_properties.tabColor = "0070C0"
    ws.cell(row=1, column=1, value="🏆 LOBBY QUALITY & LEAGUE EXPERIENCE (last 6 months)").font = \
        Font(bold=True, size=13, name="Arial")
    ws.merge_cells("A1:M1")
    ws.cell(row=2, column=1, value="Median lobby rank = the average medal of matches they queue into. "
            "Players whose lobbies sit above their own medal are undervalued. League games = Captains "
            "Mode / 10-stack practice+tournament lobbies (LD2L, AD2L, RD2L etc.), not pubs.").font = \
        Font(italic=True, name="Arial", size=9, color="666666")
    ws.merge_cells("A2:M2")

    headers = ["Player", "MMR", "Own Rank", "Median Lobby", "Gap (stars)", "Punches Up?",
               "Solo %", "Solo WR", "Party WR", "League Games", "League W%",
               "League Heroes", "Verdict"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = \
            HDR_FONT, PatternFill("solid", fgColor="003366"), HDR_ALIGN, BORDER

    def sort_key(pd):
        g = pd["data"]["punch_gap_stars"]
        return -(g if g is not None else -99)

    row = 5
    for pd in sorted(sorted_players, key=sort_key):
        p, d = pd["player"], pd["data"]
        if not d["lobby_sample"] and not d["league_matches"]:
            continue

        if d["punches_above"]:
            verdict = "🔥 UNDERVALUED - queues above own medal"
        elif d["has_league_exp"] and d["league_winrate"] >= 50:
            verdict = f"League veteran ({d['league_matches']}g, {d['league_winrate']}%)"
        elif d["has_league_exp"]:
            verdict = f"Has league exp ({d['league_matches']} games)"
        elif d["solo_wr"] is not None and d["solo_wr"] >= 55 and d["solo_n"] >= 30:
            verdict = f"Strong solo queue player ({d['solo_wr']}%)"
        else:
            verdict = ""

        gap = d["punch_gap_stars"]
        vals = [p["name"], p["mmr"], d["rank_str"], d["lobby_rank_str"],
                (f"+{gap}" if gap and gap > 0 else gap) if gap is not None else "—",
                "✅ YES" if d["punches_above"] else "",
                d["solo_pct"] if d["solo_pct"] is not None else "—",
                d["solo_wr"] if d["solo_wr"] is not None else "—",
                d["party_wr"] if d["party_wr"] is not None else "—",
                d["league_matches"] or "—",
                d["league_winrate"] if d["league_matches"] else "—",
                ", ".join(d["league_heroes"][:6]) or "—", verdict]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=_safe(v))
            cell.font = BOLD_FONT if c == 1 else DATA_FONT
            cell.border = BORDER
            cell.alignment = WRAP if c >= 12 else CENTER
            if d["punches_above"]:
                cell.fill = PatternFill("solid", fgColor="E8F5E9")
        if d["punches_above"]:
            for col in (4, 5, 6):
                ws.cell(row=row, column=col).font = GREEN_BOLD
                ws.cell(row=row, column=col).fill = GREEN_FILL
        if d["league_matches"]:
            ws.cell(row=row, column=10).fill = BLUE_FILL
        row += 1

    for c, w in {"A": 26, "B": 8, "C": 13, "D": 13, "E": 10, "F": 11, "G": 8,
                 "H": 9, "I": 9, "J": 12, "K": 10, "L": 35, "M": 40}.items():
        ws.column_dimensions[c].width = w


def _build_compare(ws):
    ws.sheet_properties.tabColor = "70AD47"
    ws.cell(row=1, column=1, value="Use this sheet to compare players side-by-side during draft "
            "(the HTML dashboard has an interactive version)").font = Font(bold=True, size=11, name="Arial")
    ws.merge_cells("A1:F1")

    headers = ["Stat", "Player A", "Player B", "Player C", "Player D", "Player E"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill = HDR_FONT, PatternFill("solid", fgColor="2D7D2D")

    stats = ["Name", "Listed MMR", "OD Rank", "MMR Check", "Est. Cost", "Last Cost",
             "Pos Prefs", "Actual Lanes",
             "Win%", "Form 30d", "Total Games", "Recent KDA", "Avg GPM", "Avg XPM",
             "Heroes Played", "Last Active (days)", "Median Lobby Rank", "Punches Up?",
             "Solo WR", "League Games", "League Win%", "Top Hero #1", "Top Hero #2",
             "Top Hero #3", "Signature Heroes", "Statement", "Your Notes"]
    for r, s in enumerate(stats, 4):
        cell = ws.cell(row=r, column=1, value=s)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="E2EFDA")

    for c, w in {"A": 18, "B": 25, "C": 25, "D": 25, "E": 25, "F": 25}.items():
        ws.column_dimensions[c].width = w


def generate_spreadsheet(players_with_data, output_file):
    wb = Workbook()
    sorted_players = sorted(players_with_data, key=lambda x: x["player"]["mmr"], reverse=True)

    ws = wb.active
    ws.title = "Scouting Board"
    _build_board(ws, sorted_players)
    _build_by_role(wb.create_sheet("By Role"), sorted_players)
    _build_value_picks(wb.create_sheet("🔥 Value Picks"), sorted_players)
    _build_signatures(wb.create_sheet("🎯 Signature Heroes"), sorted_players)
    _build_lobby_intel(wb.create_sheet("🏆 Lobby & League Intel"), sorted_players)
    _build_compare(wb.create_sheet("Head-to-Head Compare"))

    wb.save(output_file)
    print(f"\n✅ Spreadsheet saved to {output_file}")
